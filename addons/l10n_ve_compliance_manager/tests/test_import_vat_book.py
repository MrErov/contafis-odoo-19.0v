import base64
from io import BytesIO
from pathlib import Path

import openpyxl
from odoo.tests import TransactionCase, tagged


@tagged('post_install', '-at_install')
class TestImportVatBook(TransactionCase):
    """Tests para importación de Libro de Compras/Ventas (Fase G)."""

    def setUp(self):
        super().setUp()
        # Aislar BD: limpiar datos de corridas previas
        self.env['l10n.ve.vat.book.line'].search([]).unlink()
        self.env['l10n.ve.import.line'].search([]).unlink()
        self.env['l10n.ve.import.wizard'].search([]).unlink()
        self.env['l10n.ve.vat.return'].search([]).unlink()

        self.Wizard = self.env['l10n.ve.import.wizard']

    def _create_vat_book_excel(self, rows_data, book_type='purchase', sheet_name=None, header_row=3):
        """
        Genera un Excel en memoria con estructura de Libro Compras/Ventas.

        Args:
            rows_data: Lista de dicts con datos de facturas
            book_type: 'purchase' o 'sale'
            sheet_name: Nombre de la hoja (default: COMPRAS o VENTAS)
            header_row: Fila 1-based donde están los headers

        Returns:
            bytes: Contenido del archivo xlsx
        """
        wb = openpyxl.Workbook()
        ws = wb.active

        if sheet_name is None:
            ws.title = 'COMPRAS' if book_type == 'purchase' else 'VENTAS'
        else:
            ws.title = sheet_name

        # Filas previas al header (pueden tener título, período, etc.)
        if header_row > 1:
            ws.cell(row=1, column=1, value='LIBRO DE COMPRAS' if book_type == 'purchase' else 'LIBRO DE VENTAS')
            ws.cell(row=2, column=1, value='Mes AGOSTO 2026')

        # Headers según tipo
        if book_type == 'purchase':
            headers = [
                'R.I.F.',
                'Nombre o Razon Social',
                'Numero de Factura',
                'Numero de Control',
                'Fecha',
                'Base Alicuota General 16%',
                'I.V.A. Alicuota General 16%',
                'Base Alicuota Reducida',
                'I.V.A. Alicuota Reducida',
                'Compras NO SUJETAS',
                'Compras sin Derecho a Credito (Nacional)',
                'Base Importación 16%',
                'I.V.A. de importación 16%',
                'Nº Comprob. Retención 75%',
                'IVA Retenido (al Vendedor)',
                'IVA Retenido (a Terceros)',
                'Anticipo IVA (Importación)',
            ]
        else:
            headers = [
                'R.I.F',
                'Nombre o Razon Social',
                'Numero de Factura 0 reporte Z',
                'Numero de Control',
                'Fecha de la Factura',
                'Ventas internas No sujetas',
                'Ventas internas no gravadas (No Contrib)',
                'Base Imponible (No Contrib)',
                'Impuesto IVA (No Contrib)',
                'Base Imponible (Contrib)',
                'Impuesto IVA (Contrib)',
                'Nº Comprob. Retención 75% IVA',
                'Iva Retenido (por comprador)',
            ]

        # Escribir headers en la fila indicada
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=header_row, column=col_idx, value=header)

        # Datos
        for row_idx, row_data in enumerate(rows_data, header_row + 1):
            if book_type == 'purchase':
                ws.cell(row=row_idx, column=1, value=row_data.get('partner_vat'))
                ws.cell(row=row_idx, column=2, value=row_data.get('partner_name'))
                ws.cell(row=row_idx, column=3, value=row_data.get('invoice_number'))
                ws.cell(row=row_idx, column=4, value=row_data.get('control_number'))
                ws.cell(row=row_idx, column=5, value=row_data.get('invoice_date'))
                ws.cell(row=row_idx, column=6, value=row_data.get('base_general'))
                ws.cell(row=row_idx, column=7, value=row_data.get('vat_general'))
                ws.cell(row=row_idx, column=8, value=row_data.get('base_reduced'))
                ws.cell(row=row_idx, column=9, value=row_data.get('vat_reduced'))
                ws.cell(row=row_idx, column=10, value=row_data.get('base_not_subject'))
                ws.cell(row=row_idx, column=11, value=row_data.get('base_no_credit'))
                ws.cell(row=row_idx, column=12, value=row_data.get('base_import_16'))
                ws.cell(row=row_idx, column=13, value=row_data.get('vat_import_16'))
                ws.cell(row=row_idx, column=14, value=row_data.get('retention_number'))
                ws.cell(row=row_idx, column=15, value=row_data.get('vat_retained_vendor'))
                ws.cell(row=row_idx, column=16, value=row_data.get('vat_retained_third'))
                ws.cell(row=row_idx, column=17, value=row_data.get('anticipo_import'))
            else:
                ws.cell(row=row_idx, column=1, value=row_data.get('partner_vat'))
                ws.cell(row=row_idx, column=2, value=row_data.get('partner_name'))
                ws.cell(row=row_idx, column=3, value=row_data.get('invoice_number'))
                ws.cell(row=row_idx, column=4, value=row_data.get('control_number'))
                ws.cell(row=row_idx, column=5, value=row_data.get('invoice_date'))
                ws.cell(row=row_idx, column=6, value=row_data.get('base_not_subject'))
                ws.cell(row=row_idx, column=7, value=row_data.get('base_not_taxed'))
                ws.cell(row=row_idx, column=8, value=row_data.get('base_general_non_contrib'))
                ws.cell(row=row_idx, column=9, value=row_data.get('vat_general_non_contrib'))
                ws.cell(row=row_idx, column=10, value=row_data.get('base_general_contrib'))
                ws.cell(row=row_idx, column=11, value=row_data.get('vat_general_contrib'))
                ws.cell(row=row_idx, column=12, value=row_data.get('retention_number'))
                ws.cell(row=row_idx, column=13, value=row_data.get('vat_retained_buyer'))

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def _create_vat_book_both_excel(self, purchase_rows, sale_rows, header_row=3):
        """
        Genera un Excel con 2 hojas: COMPRAS y VENTAS.

        Args:
            purchase_rows: Lista de dicts para COMPRAS
            sale_rows: Lista de dicts para VENTAS
            header_row: Fila 1-based donde están los headers

        Returns:
            bytes: Contenido del archivo xlsx
        """
        wb = openpyxl.Workbook()

        # Hoja COMPRAS
        ws_compras = wb.active
        ws_compras.title = 'COMPRAS'

        if header_row > 1:
            ws_compras.cell(row=1, column=1, value='LIBRO DE COMPRAS')
            ws_compras.cell(row=2, column=1, value='Mes AGOSTO 2026')

        purchase_headers = [
            'R.I.F.',
            'Nombre o Razon Social',
            'Numero de Factura',
            'Numero de Control',
            'Fecha',
            'Base Alicuota General 16%',
            'I.V.A. Alicuota General 16%',
            'Base Alicuota Reducida',
            'I.V.A. Alicuota Reducida',
            'Compras NO SUJETAS',
            'Compras sin Derecho a Credito (Nacional)',
            'Base Importación 16%',
            'I.V.A. de importación 16%',
            'Nº Comprob. Retención 75%',
            'IVA Retenido (al Vendedor)',
            'IVA Retenido (a Terceros)',
            'Anticipo IVA (Importación)',
        ]

        for col_idx, header in enumerate(purchase_headers, 1):
            ws_compras.cell(row=header_row, column=col_idx, value=header)

        for row_idx, row_data in enumerate(purchase_rows, header_row + 1):
            ws_compras.cell(row=row_idx, column=1, value=row_data.get('partner_vat'))
            ws_compras.cell(row=row_idx, column=2, value=row_data.get('partner_name'))
            ws_compras.cell(row=row_idx, column=3, value=row_data.get('invoice_number'))
            ws_compras.cell(row=row_idx, column=4, value=row_data.get('control_number'))
            ws_compras.cell(row=row_idx, column=5, value=row_data.get('invoice_date'))
            ws_compras.cell(row=row_idx, column=6, value=row_data.get('base_general'))
            ws_compras.cell(row=row_idx, column=7, value=row_data.get('vat_general'))
            ws_compras.cell(row=row_idx, column=8, value=row_data.get('base_reduced'))
            ws_compras.cell(row=row_idx, column=9, value=row_data.get('vat_reduced'))
            ws_compras.cell(row=row_idx, column=10, value=row_data.get('base_not_subject'))
            ws_compras.cell(row=row_idx, column=11, value=row_data.get('base_no_credit'))
            ws_compras.cell(row=row_idx, column=12, value=row_data.get('base_import_16'))
            ws_compras.cell(row=row_idx, column=13, value=row_data.get('vat_import_16'))
            ws_compras.cell(row=row_idx, column=14, value=row_data.get('retention_number'))
            ws_compras.cell(row=row_idx, column=15, value=row_data.get('vat_retained_vendor'))
            ws_compras.cell(row=row_idx, column=16, value=row_data.get('vat_retained_third'))
            ws_compras.cell(row=row_idx, column=17, value=row_data.get('anticipo_import'))

        # Hoja VENTAS
        ws_ventas = wb.create_sheet('VENTAS')

        if header_row > 1:
            ws_ventas.cell(row=1, column=1, value='LIBRO DE VENTAS')
            ws_ventas.cell(row=2, column=1, value='Mes AGOSTO 2026')

        sale_headers = [
            'R.I.F',
            'Nombre o Razon Social',
            'Numero de Factura 0 reporte Z',
            'Numero de Control',
            'Fecha de la Factura',
            'Ventas internas No sujetas',
            'Ventas internas no gravadas (No Contrib)',
            'Base Imponible (No Contrib)',
            'Impuesto IVA (No Contrib)',
            'Base Imponible (Contrib)',
            'Impuesto IVA (Contrib)',
            'Nº Comprob. Retención 75% IVA',
            'Iva Retenido (por comprador)',
        ]

        for col_idx, header in enumerate(sale_headers, 1):
            ws_ventas.cell(row=header_row, column=col_idx, value=header)

        for row_idx, row_data in enumerate(sale_rows, header_row + 1):
            ws_ventas.cell(row=row_idx, column=1, value=row_data.get('partner_vat'))
            ws_ventas.cell(row=row_idx, column=2, value=row_data.get('partner_name'))
            ws_ventas.cell(row=row_idx, column=3, value=row_data.get('invoice_number'))
            ws_ventas.cell(row=row_idx, column=4, value=row_data.get('control_number'))
            ws_ventas.cell(row=row_idx, column=5, value=row_data.get('invoice_date'))
            ws_ventas.cell(row=row_idx, column=6, value=row_data.get('base_not_subject'))
            ws_ventas.cell(row=row_idx, column=7, value=row_data.get('base_not_taxed'))
            ws_ventas.cell(row=row_idx, column=8, value=row_data.get('base_general_non_contrib'))
            ws_ventas.cell(row=row_idx, column=9, value=row_data.get('vat_general_non_contrib'))
            ws_ventas.cell(row=row_idx, column=10, value=row_data.get('base_general_contrib'))
            ws_ventas.cell(row=row_idx, column=11, value=row_data.get('vat_general_contrib'))
            ws_ventas.cell(row=row_idx, column=12, value=row_data.get('retention_number'))
            ws_ventas.cell(row=row_idx, column=13, value=row_data.get('vat_retained_buyer'))

        output = BytesIO()
        wb.save(output)
        return output.getvalue()
        wb.save(output)
        return output.getvalue()

    def _create_wizard(self, import_type='vat_book_purchase'):
        """Helper para crear wizard con datos por defecto."""
        return self.Wizard.create({
            'import_type': import_type,
        })

    def test_import_types_exist(self):
        """Verifica que los nuevos import_type están en la selección."""
        wizard = self.env['l10n.ve.import.wizard']
        types = dict(wizard._fields['import_type'].selection)
        self.assertIn('vat_book_purchase', types)
        self.assertIn('vat_book_sale', types)
        self.assertIn('vat_book_both', types)
        self.assertEqual(types['vat_book_purchase'], 'Libro de Compras (IVA)')
        self.assertEqual(types['vat_book_sale'], 'Libro de Ventas (IVA)')
        self.assertEqual(types['vat_book_both'], 'Libro de Compras y Ventas (IVA)')

    def test_wizard_has_vat_book_fields(self):
        """Verifica que el wizard tiene los campos period_month y sheet_name."""
        wizard = self.Wizard.create({
            'import_type': 'vat_book_purchase',
        })
        # Los campos deben existir en el modelo
        self.assertTrue('period_month' in wizard._fields)
        self.assertTrue('sheet_name' in wizard._fields)
        # period_month es Char
        self.assertEqual(wizard._fields['period_month'].type, 'char')
        # sheet_name es Char readonly
        self.assertEqual(wizard._fields['sheet_name'].type, 'char')
        self.assertTrue(wizard._fields['sheet_name'].readonly)

    def test_detect_vat_book_sheet(self):
        """Verifica detección de hoja COMPRAS/VENTAS."""
        wizard = self._create_wizard('vat_book_purchase')

        # Crear Excel con 2 hojas
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = 'COMPRAS'
        ws2 = wb.create_sheet('VENTAS')

        output = BytesIO()
        wb.save(output)
        excel_content = output.getvalue()

        # Test purchase
        wb_test = openpyxl.load_workbook(BytesIO(excel_content), read_only=True)
        sheet = wizard._detect_vat_book_sheet(wb_test, 'purchase')
        self.assertEqual(sheet, 'COMPRAS')

        # Test sale
        wb_test = openpyxl.load_workbook(BytesIO(excel_content), read_only=True)
        sheet = wizard._detect_vat_book_sheet(wb_test, 'sale')
        self.assertEqual(sheet, 'VENTAS')

        # Test fallback purchase (hoja 0)
        wb2 = openpyxl.Workbook()
        wb2.active.title = 'Hoja1'
        wb2.create_sheet('Hoja2')
        output2 = BytesIO()
        wb2.save(output2)
        wb_test2 = openpyxl.load_workbook(BytesIO(output2.getvalue()), read_only=True)
        sheet = wizard._detect_vat_book_sheet(wb_test2, 'purchase')
        self.assertEqual(sheet, 'Hoja1')

        # Test fallback sale (hoja 1)
        wb_test2 = openpyxl.load_workbook(BytesIO(output2.getvalue()), read_only=True)
        sheet = wizard._detect_vat_book_sheet(wb_test2, 'sale')
        self.assertEqual(sheet, 'Hoja2')

    def test_parse_vat_book_headers(self):
        """Verifica detección de fila header y mapeo de columnas."""
        wizard = self._create_wizard('vat_book_purchase')

        # Crear Excel con header en fila 7
        rows_data = [
            {
                'partner_vat': 'J-31527189-4',
                'partner_name': 'Empresa Test',
                'invoice_number': '001',
                'control_number': '001',
                'invoice_date': '15/08/2026',
                'base_general': 1000,
                'vat_general': 160,
                'base_reduced': 0,
                'vat_reduced': 0,
                'base_not_subject': 0,
                'base_no_credit': 0,
                'base_import_16': 0,
                'vat_import_16': 0,
                'retention_number': '',
                'vat_retained_vendor': 0,
                'vat_retained_third': 0,
                'anticipo_import': 0,
            }
        ]

        excel_content = self._create_vat_book_excel(rows_data, book_type='purchase', header_row=7)
        excel_b64 = base64.b64encode(excel_content)

        wizard.write({
            'file': excel_b64,
            'filename': 'test_vat_book.xlsx',
        })

        # Cargar archivo (esto llama _parse_vat_book_excel internamente)
        wizard.action_load_file()

        # Verificar que se detectó la hoja
        self.assertEqual(wizard.sheet_name, 'COMPRAS')

        # Verificar que se crearon líneas de preview
        self.assertEqual(len(wizard.line_ids), 1)

        # Verificar datos de la línea
        line = wizard.line_ids[0]
        self.assertEqual(line.data['partner_vat'], 'J-31527189-4')
        self.assertEqual(line.data['partner_name'], 'Empresa Test')
        self.assertEqual(line.data['invoice_number'], '001')
        self.assertEqual(line.data['base_general'], 1000.0)
        self.assertEqual(line.data['vat_general'], 160.0)

        # Verificar period_month detectado
        self.assertEqual(wizard.period_month, '2026-08')

    def test_import_vat_book_purchase_creates_line(self):
        """Importa COMPRAS base_general > 0 → vat.book.line operation_code='33'."""
        wizard = self._create_wizard('vat_book_purchase')
        rows = [{
            'partner_vat': 'J-31527189-4',
            'partner_name': 'Proveedor Test',
            'invoice_number': '001',
            'control_number': '001',
            'invoice_date': '15/08/2026',
            'base_general': 1000,
            'vat_general': 160,
            'base_reduced': 0,
            'vat_reduced': 0,
            'base_not_subject': 0,
            'base_no_credit': 0,
            'base_import_16': 0,
            'vat_import_16': 0,
            'retention_number': '',
            'vat_retained_vendor': 0,
            'vat_retained_third': 0,
            'anticipo_import': 0,
        }]
        excel = self._create_vat_book_excel(rows, book_type='purchase')
        wizard.write({'file': base64.b64encode(excel), 'filename': 'test.xlsx'})
        wizard.action_load_file()
        wizard.action_import()

        lines = self.env['l10n.ve.vat.book.line'].search([
            ('invoice_number', '=', '001'),
            ('period_month', '=', '2026-08'),
        ])
        self.assertEqual(len(lines), 1)
        self.assertEqual(lines.operation_code, '33')
        self.assertEqual(lines.base_general, 1000)
        self.assertEqual(lines.vat_general, 160)
        self.assertEqual(lines.retention_direction, 'to_vendor')

    def test_import_vat_book_sale_creates_line(self):
        """Importa VENTAS base_general_contrib > 0 → vat.book.line operation_code='42'."""
        wizard = self._create_wizard('vat_book_sale')
        wizard.period_month = '2026-08'  # Fallback explícito
        rows = [{
            'partner_vat': 'V-12345678-5',
            'partner_name': 'Cliente Test',
            'invoice_number': '001-SALE',
            'control_number': '001',
            'invoice_date': '15/08/2026',
            'base_not_subject': 0,
            'base_not_taxed': 0,
            'base_general_non_contrib': 0,
            'vat_general_non_contrib': 0,
            'base_general_contrib': 1000,
            'vat_general_contrib': 160,
            'retention_number': '',
            'vat_retained_buyer': 50,
        }]
        excel = self._create_vat_book_excel(rows, book_type='sale')
        wizard.write({'file': base64.b64encode(excel), 'filename': 'test.xlsx'})
        wizard.action_load_file()
        wizard.action_import()

        lines = self.env['l10n.ve.vat.book.line'].search([
            ('invoice_number', '=', '001-SALE'),
            ('period_month', '=', '2026-08'),
        ])
        self.assertEqual(len(lines), 1)
        self.assertEqual(lines.operation_code, '42')
        self.assertEqual(lines.base_general, 1000)
        self.assertEqual(lines.vat_general, 160)
        self.assertEqual(lines.retention_direction, 'by_buyer')
        self.assertEqual(lines.vat_retained, 50)

    def test_import_vat_book_split_retention(self):
        """Compra con vat_retained_vendor > 0 Y vat_retained_third > 0 → 2 líneas."""
        wizard = self._create_wizard('vat_book_purchase')
        rows = [{
            'partner_vat': 'J-31527189-4',
            'partner_name': 'Proveedor Split',
            'invoice_number': '002',
            'control_number': '002',
            'invoice_date': '20/08/2026',
            'base_general': 2000,
            'vat_general': 320,
            'base_reduced': 0,
            'vat_reduced': 0,
            'base_not_subject': 0,
            'base_no_credit': 0,
            'base_import_16': 0,
            'vat_import_16': 0,
            'retention_number': 'RET-001',
            'vat_retained_vendor': 100,
            'vat_retained_third': 50,
            'anticipo_import': 0,
        }]
        excel = self._create_vat_book_excel(rows, book_type='purchase')
        wizard.write({'file': base64.b64encode(excel), 'filename': 'test.xlsx'})
        wizard.action_load_file()
        wizard.action_import()

        # Deben crearse 2 líneas: una to_vendor, una to_third
        lines = self.env['l10n.ve.vat.book.line'].search([
            ('invoice_number', '=', '002'),
            ('period_month', '=', '2026-08'),
        ])
        self.assertEqual(len(lines), 2, 'Debe crear 2 líneas por split de retención')

        # Verificar dirección to_vendor
        line_vendor = lines.filtered(lambda l: l.retention_direction == 'to_vendor')
        self.assertEqual(len(line_vendor), 1)
        self.assertEqual(line_vendor.vat_retained, 100)
        self.assertEqual(line_vendor.operation_code, '33')

        # Verificar dirección to_third
        line_third = lines.filtered(lambda l: l.retention_direction == 'to_third')
        self.assertEqual(len(line_third), 1)
        self.assertEqual(line_third.vat_retained, 50)
        self.assertEqual(line_third.operation_code, '33')

        # Ambas con mismo operation_code pero distinta direction → no colisionan
        self.assertNotEqual(line_vendor.id, line_third.id)

    def test_download_purchase_template(self):
        """Verifica que Descargar Plantilla funciona para vat_book_purchase."""
        wizard = self._create_wizard('vat_book_purchase')
        result = wizard.action_download_template()

        self.assertEqual(result['type'], 'ir.actions.act_url')
        self.assertIn('plantilla_vat_book_purchase.xlsx', result['url'])
        self.assertTrue(wizard.template_file)

        import base64
        xlsx = base64.b64decode(wizard.template_file)
        wb = openpyxl.load_workbook(BytesIO(xlsx))
        ws = wb.active

        # Verificar headers en la primera fila (plantilla)
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        expected = [
            'R.I.F.', 'Nombre o Razon Social', 'Tipo Doc', 'Numero de Factura',
            'Numero de Control', 'Fecha', 'Base Alicuota General 16%',
            'I.V.A. Alicuota General 16%', 'Base Alicuota Reducida',
            'I.V.A. Alicuota Reducida', 'Compras NO SUJETAS',
            'Compras sin Derecho a Credito (Nacional)', 'Base Importación 16%',
            'I.V.A. de importación 16%', 'Nº Comprob. Retención 75%',
            'IVA Retenido (al Vendedor)', 'IVA Retenido (a Terceros)',
            'Anticipo IVA (Importación)',
        ]
        self.assertEqual(headers, expected)

        # Verificar fila de ejemplo (fila 2 existe, vacía)
        self.assertEqual(ws.max_row, 2)

    def test_download_sale_template(self):
        """Verifica que Descargar Plantilla funciona para vat_book_sale."""
        wizard = self._create_wizard('vat_book_sale')
        result = wizard.action_download_template()

        self.assertEqual(result['type'], 'ir.actions.act_url')
        self.assertIn('plantilla_vat_book_sale.xlsx', result['url'])
        self.assertTrue(wizard.template_file)

        import base64
        xlsx = base64.b64decode(wizard.template_file)
        wb = openpyxl.load_workbook(BytesIO(xlsx))
        ws = wb.active

        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        expected = [
            'R.I.F.', 'Nombre o Razon Social', 'Tipo Doc', 'Numero de Factura',
            'Numero de Control', 'Fecha', 'Ventas internas No sujetas',
            'Ventas internas no gravadas (No Contrib)', 'Base Imponible (No Contrib)',
            'Impuesto IVA (No Contrib)', 'Base Imponible (Contrib)',
            'Impuesto IVA (Contrib)', 'Nº Comprob. Retención 75% IVA',
            'Iva Retenido (por comprador)',
        ]
        self.assertEqual(headers, expected)

        # Verificar fila de ejemplo (fila 2 existe, vacía)
        self.assertEqual(ws.max_row, 2)

    def test_import_vat_book_both_loads_two_sheets(self):
        """vat_book_both carga ambas hojas y marca _book_type en cada línea."""
        wizard = self._create_wizard('vat_book_both')

        purchase_rows = [{
            'partner_vat': 'J-31527189-4',
            'partner_name': 'Proveedor Test',
            'invoice_number': '001',
            'control_number': '001',
            'invoice_date': '15/08/2026',
            'base_general': 1000,
            'vat_general': 160,
            'base_reduced': 0,
            'vat_reduced': 0,
            'base_not_subject': 0,
            'base_no_credit': 500,
            'base_import_16': 0,
            'vat_import_16': 0,
            'retention_number': '',
            'vat_retained_vendor': 0,
            'vat_retained_third': 0,
            'anticipo_import': 0,
        }]
        sale_rows = [{
            'partner_vat': 'V-12345678-5',
            'partner_name': 'Cliente Test',
            'invoice_number': '001-SALE',
            'control_number': '001',
            'invoice_date': '15/08/2026',
            'base_not_subject': 0,
            'base_not_taxed': 0,
            'base_general_non_contrib': 0,
            'vat_general_non_contrib': 0,
            'base_general_contrib': 1000,
            'vat_general_contrib': 160,
            'retention_number': '',
            'vat_retained_buyer': 50,
        }]

        excel = self._create_vat_book_both_excel(purchase_rows, sale_rows)
        wizard.write({'file': base64.b64encode(excel), 'filename': 'test.xlsx'})
        wizard.action_load_file()

        # Verificar que se crearon líneas de ambas hojas
        self.assertTrue(len(wizard.line_ids) >= 2)

        # Verificar _book_type
        purchase_lines = wizard.line_ids.filtered(lambda l: l.data.get('_book_type') == 'purchase')
        sale_lines = wizard.line_ids.filtered(lambda l: l.data.get('_book_type') == 'sale')

        self.assertTrue(len(purchase_lines) >= 1, 'Debe haber líneas purchase')
        self.assertTrue(len(sale_lines) >= 1, 'Debe haber líneas sale')

        # Verificar multi-rate en purchase (base_general + base_no_credit = 2 líneas)
        self.assertEqual(len(purchase_lines), 2)

    def test_import_vat_book_both_imports_all(self):
        """vat_book_both importa creando vat.book.line de ambos tipos."""
        wizard = self._create_wizard('vat_book_both')
        wizard.period_month = '2026-08'

        # 1 factura compra con base_general (1 línea)
        # 1 factura compra con base_general + base_no_credit (2 líneas = multi-rate)
        purchase_rows = [
            {
                'partner_vat': 'J-31527189-4',
                'partner_name': 'Proveedor 1',
                'invoice_number': '001',
                'control_number': '001',
                'invoice_date': '15/08/2026',
                'base_general': 1000,
                'vat_general': 160,
                'base_reduced': 0,
                'vat_reduced': 0,
                'base_not_subject': 0,
                'base_no_credit': 0,
                'base_import_16': 0,
                'vat_import_16': 0,
                'retention_number': '',
                'vat_retained_vendor': 0,
                'vat_retained_third': 0,
                'anticipo_import': 0,
            },
            {
                'partner_vat': 'J-39204516-4',
                'partner_name': 'Proveedor 2',
                'invoice_number': '002',
                'control_number': '002',
                'invoice_date': '20/08/2026',
                'base_general': 2000,
                'vat_general': 320,
                'base_reduced': 0,
                'vat_reduced': 0,
                'base_not_subject': 0,
                'base_no_credit': 500,  # multi-rate: genera 2 líneas
                'base_import_16': 0,
                'vat_import_16': 0,
                'retention_number': '',
                'vat_retained_vendor': 0,
                'vat_retained_third': 0,
                'anticipo_import': 0,
            },
        ]
        # 2 facturas venta (1 contrib, 1 non_contrib)
        sale_rows = [
            {
                'partner_vat': 'V-12345678-5',
                'partner_name': 'Cliente 1',
                'invoice_number': '001-SALE',
                'control_number': '001',
                'invoice_date': '15/08/2026',
                'base_not_subject': 0,
                'base_not_taxed': 0,
                'base_general_non_contrib': 0,
                'vat_general_non_contrib': 0,
                'base_general_contrib': 1000,
                'vat_general_contrib': 160,
                'retention_number': '',
                'vat_retained_buyer': 50,
            },
            {
                'partner_vat': 'V-26920712-3',
                'partner_name': 'Cliente 2',
                'invoice_number': '002-SALE',
                'control_number': '002',
                'invoice_date': '20/08/2026',
                'base_not_subject': 0,
                'base_not_taxed': 0,
                'base_general_non_contrib': 1000,
                'vat_general_non_contrib': 160,
                'base_general_contrib': 0,
                'vat_general_contrib': 0,
                'retention_number': '',
                'vat_retained_buyer': 0,
            },
        ]

        excel = self._create_vat_book_both_excel(purchase_rows, sale_rows)
        wizard.write({'file': base64.b64encode(excel), 'filename': 'test.xlsx'})
        wizard.action_load_file()
        wizard.action_import()

        # Verificar vat.book.line creados
        vbl = self.env['l10n.ve.vat.book.line'].search([
            ('period_month', '=', '2026-08'),
        ])

        purchase_vbl = vbl.filtered(lambda l: l.book_type == 'purchase')
        sale_vbl = vbl.filtered(lambda l: l.book_type == 'sale')

        # purchase: 2 facturas, una con multi-rate = 3 líneas
        self.assertEqual(len(purchase_vbl), 3, 'Debe crear 3 líneas purchase (2 facturas, 1 multi-rate)')

        # sale: 2 facturas = 2 líneas
        self.assertEqual(len(sale_vbl), 2, 'Debe crear 2 líneas sale')

        # Verificar operation_codes
        # purchase: general=33, no_credit=30
        self.assertTrue(all(l.operation_code in ('33', '30') for l in purchase_vbl))
        self.assertTrue(all(l.operation_code in ('42', '443') for l in sale_vbl))

    def _compute_expected_totals_from_fixture(self, book_type):
        """Suma los totales del Excel real para validar contra el import.

        NOTA: El import combina base_import_16 en base_general y vat_import_16 en vat_general
        (ver _build_vat_book_vals en import_wizard.py). Este helper replica esa lógica
        para que la comparación sea apple-to-apple.
        """
        from pathlib import Path

        import openpyxl

        fixture_path = Path(__file__).parent / 'fixtures' / 'Libro_COMPRAS_VENTAS_Ficticio.xlsx'
        wb = openpyxl.load_workbook(fixture_path, data_only=True)

        if book_type == 'purchase':
            ws = wb['COMPRAS']
            data_start, data_end = 8, 17

            base_general = 0.0
            vat_general = 0.0
            base_no_credit = 0.0
            base_import_16 = 0.0
            vat_import_16 = 0.0
            vat_retained_vendor = 0.0

            for r in range(data_start, data_end + 1):
                v = ws.cell(row=r, column=23).value
                if isinstance(v, (int, float)):
                    base_general += v
                v = ws.cell(row=r, column=24).value
                if isinstance(v, (int, float)):
                    vat_general += v
                v = ws.cell(row=r, column=21).value
                if isinstance(v, (int, float)):
                    base_no_credit += v
                v = ws.cell(row=r, column=19).value
                if isinstance(v, (int, float)):
                    base_import_16 += v
                v = ws.cell(row=r, column=20).value
                if isinstance(v, (int, float)):
                    vat_import_16 += v
                v = ws.cell(row=r, column=29).value
                if isinstance(v, (int, float)):
                    vat_retained_vendor += v

            # Replicar lógica del import: combinar import en general
            return {
                'base_general': base_general + base_import_16,
                'vat_general': vat_general + vat_import_16,
                'base_no_credit': base_no_credit,
                'base_import_16': base_import_16,
                'vat_import_16': vat_import_16,
                'vat_retained_vendor': vat_retained_vendor,
            }
        else:
            ws = wb['VENTAS']
            data_start, data_end = 12, 21
            cols = {
                'base_general_contrib': 22,
                'vat_general_contrib': 24,
                'base_general_non_contrib': 18,
                'vat_general_non_contrib': 20,
                'vat_retained_buyer': 27,
            }

            totals = {k: 0.0 for k in cols}
            for r in range(data_start, data_end + 1):
                for key, col in cols.items():
                    v = ws.cell(row=r, column=col).value
                    if isinstance(v, (int, float)):
                        totals[key] += v
            return totals

    def test_import_real_fixture_purchase_12_lines(self):
        """Fixture real COMPRAS: 10 facturas + 2 multi-rate = 12 import.line."""
        wizard = self._create_wizard('vat_book_purchase')

        fixture_path = Path(__file__).parent / 'fixtures' / 'Libro_COMPRAS_VENTAS_Ficticio.xlsx'
        wizard.write({
            'file': base64.b64encode(fixture_path.read_bytes()),
            'filename': fixture_path.name,
        })
        wizard.action_load_file()

        # 10 facturas: 8 normales (1 línea c/u) + 2 con multi-rate (2 líneas c/u) = 12
        self.assertEqual(len(wizard.line_ids), 12,
            f'Esperado 12 import.line, got {len(wizard.line_ids)}')

        # Todas deben ser purchase
        for line in wizard.line_ids:
            bt = line.data.get('_book_type')
            self.assertIn(bt, ('purchase', ''), f'Línea con _book_type inesperado: {bt}')

    def test_import_real_fixture_sale_9_lines(self):
        """Fixture real VENTAS: 10 facturas - 1 anulada = 9 import.line."""
        wizard = self._create_wizard('vat_book_sale')

        fixture_path = Path(__file__).parent / 'fixtures' / 'Libro_COMPRAS_VENTAS_Ficticio.xlsx'
        wizard.write({
            'file': base64.b64encode(fixture_path.read_bytes()),
            'filename': fixture_path.name,
        })
        wizard.action_load_file()

        # 10 facturas - 1 anulada = 9
        self.assertEqual(len(wizard.line_ids), 9,
            f'Esperado 9 import.line, got {len(wizard.line_ids)}')

    def test_import_real_fixture_both_creates_21_lines(self):
        """Fixture real BOTH: 12 purchase + 9 sale = 21 import.line → 21 vat.book.line."""
        wizard = self._create_wizard('vat_book_both')
        wizard.period_month = '2026-08'

        fixture_path = Path(__file__).parent / 'fixtures' / 'Libro_COMPRAS_VENTAS_Ficticio.xlsx'
        wizard.write({
            'file': base64.b64encode(fixture_path.read_bytes()),
            'filename': fixture_path.name,
        })
        wizard.action_load_file()

        # 12 + 9 = 21 import.lines
        self.assertEqual(len(wizard.line_ids), 21,
            f'Esperado 21 import.line, got {len(wizard.line_ids)}')

        purchase_lines = wizard.line_ids.filtered(lambda l: l.data.get('_book_type') == 'purchase')
        sale_lines = wizard.line_ids.filtered(lambda l: l.data.get('_book_type') == 'sale')
        self.assertEqual(len(purchase_lines), 12)
        self.assertEqual(len(sale_lines), 9)

        # Importar y verificar vat.book.line
        wizard.action_import()

        vbl = self.env['l10n.ve.vat.book.line'].search([
            ('period_month', '=', '2026-08'),
        ])

        purchase_vbl = vbl.filtered(lambda l: l.book_type == 'purchase')
        sale_vbl = vbl.filtered(lambda l: l.book_type == 'sale')

        # 12 purchase + 9 sale = 21
        self.assertEqual(len(purchase_vbl), 12)
        self.assertEqual(len(sale_vbl), 9)
        self.assertEqual(len(vbl), 21)

    def test_import_real_fixture_purchase_idempotent(self):
        """Re-importar el mismo fixture no duplica ni cambia totales."""
        wizard = self._create_wizard('vat_book_purchase')
        wizard.period_month = '2026-08'
        fixture_path = Path(__file__).parent / 'fixtures' / 'Libro_COMPRAS_VENTAS_Ficticio.xlsx'
        wizard.write({
            'file': base64.b64encode(fixture_path.read_bytes()),
            'filename': fixture_path.name,
        })
        wizard.action_load_file()
        wizard.action_import()

        vbl = self.env['l10n.ve.vat.book.line'].search([
            ('period_month', '=', '2026-08'),
            ('book_type', '=', 'purchase'),
        ])

        total_base = sum(vbl.mapped('base_general'))
        total_vat = sum(vbl.mapped('vat_general'))

        # Verificar que re-importar produce los mismos totales (upsert idempotente)
        wizard2 = self._create_wizard('vat_book_purchase')
        wizard2.period_month = '2026-08'
        wizard2.write({
            'file': base64.b64encode(fixture_path.read_bytes()),
            'filename': fixture_path.name,
        })
        wizard2.action_load_file()
        wizard2.action_import()

        vbl2 = self.env['l10n.ve.vat.book.line'].search([
            ('period_month', '=', '2026-08'),
            ('book_type', '=', 'purchase'),
        ])

        total_base2 = sum(vbl2.mapped('base_general'))
        total_vat2 = sum(vbl2.mapped('vat_general'))

        self.assertEqual(total_base, total_base2, 'Re-importar debe dar mismo base_general')
        self.assertEqual(total_vat, total_vat2, 'Re-importar debe dar mismo vat_general')

        # Verificar que los totales son positivos y razonables
        self.assertGreater(total_base, 0)

    def test_import_real_fixture_purchase_totals_matches_excel(self):
        """Verifica que los totales importados coinciden con valores
        concretos del fixture real.

        Los valores hardcodeados documentan el total del fixture
        estático. Si el fixture cambia, este test falla hasta que
        se actualicen los números.
        """
        wizard = self._create_wizard('vat_book_purchase')
        wizard.period_month = '2026-08'
        fixture_path = Path(__file__).parent / 'fixtures' / 'Libro_COMPRAS_VENTAS_Ficticio.xlsx'
        wizard.write({
            'file': base64.b64encode(fixture_path.read_bytes()),
            'filename': fixture_path.name,
        })
        wizard.action_load_file()
        wizard.action_import()

        vbl = self.env['l10n.ve.vat.book.line'].search([
            ('period_month', '=', '2026-08'),
            ('book_type', '=', 'purchase'),
        ])

        total_base = sum(vbl.mapped('base_general'))
        total_vat = sum(vbl.mapped('vat_general'))

        # Valores del fixture COMPRAS regenerado
        # (base_general + base_import_16 combinados según _build_vat_book_vals)
        # base_general = sum(col23 filas 8-17) + col19 fila 13
        # vat_general  = sum(col24 filas 8-17) + col20 fila 13
        EXPECTED_BASE_GENERAL = 2477517.43  # actualizar si el fixture cambia
        EXPECTED_VAT_GENERAL = 396402.80     # actualizar si el fixture cambia

        self.assertAlmostEqual(total_base, EXPECTED_BASE_GENERAL, places=2,
            msg=f"base_general: import={total_base} expected={EXPECTED_BASE_GENERAL}")
        self.assertAlmostEqual(total_vat, EXPECTED_VAT_GENERAL, places=2,
            msg=f"vat_general: import={total_vat} expected={EXPECTED_VAT_GENERAL}")

    def test_import_real_fixture_sale_anulada_skipped(self):
        """Factura anulada VENTAS (fila 17, invoice 000139) NO está en vat.book.line."""
        wizard = self._create_wizard('vat_book_sale')
        wizard.period_month = '2026-08'

        fixture_path = Path(__file__).parent / 'fixtures' / 'Libro_COMPRAS_VENTAS_Ficticio.xlsx'
        wizard.write({
            'file': base64.b64encode(fixture_path.read_bytes()),
            'filename': fixture_path.name,
        })
        wizard.action_load_file()
        wizard.action_import()

        # Verificar que NO existe línea con invoice_number 000139
        anulada = self.env['l10n.ve.vat.book.line'].search([
            ('invoice_number', '=', '000139'),
            ('period_month', '=', '2026-08'),
            ('book_type', '=', 'sale'),
        ])
        self.assertEqual(len(anulada), 0, 'Factura anulada 000139 no debe importarse')
