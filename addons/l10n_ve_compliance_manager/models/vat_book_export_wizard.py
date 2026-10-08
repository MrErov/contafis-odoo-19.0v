import base64
from io import BytesIO

from odoo import fields, models
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class VatBookExportWizard(models.TransientModel):
    _name = 'l10n.ve.vat.book.export.wizard'
    _description = 'Wizard de Exportación Libro de Compras/Ventas a Excel'

    period_month = fields.Char(
        string='Período (YYYY-MM)',
        required=True,
        help='Formato: 2026-08',
    )

    company_id = fields.Many2one(
        'res.company',
        string='Compañía',
        default=lambda self: self.env.company,
        required=True,
    )

    book_type = fields.Selection([
        ('purchase', 'Compras'),
        ('sale', 'Ventas'),
        ('both', 'Compras y Ventas'),
    ], string='Tipo de Libro', required=True, default='purchase')

    template_file = fields.Binary(string='Archivo Generado')
    filename = fields.Char(string='Nombre Archivo')

    def action_export_xlsx(self):
        """Genera y descarga el Excel según book_type."""
        self.ensure_one()

        if self.book_type == 'both':
            xlsx_content = self._generate_vat_book_both_xlsx()
            filename = f'Libro_Compras_Ventas_{self.period_month}.xlsx'
        elif self.book_type == 'purchase':
            xlsx_content = self._generate_vat_book_xlsx('purchase')
            filename = f'Libro_Compras_{self.period_month}.xlsx'
        else:
            xlsx_content = self._generate_vat_book_xlsx('sale')
            filename = f'Libro_Ventas_{self.period_month}.xlsx'

        self.write({
            'template_file': base64.b64encode(xlsx_content),
            'filename': filename,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/?model=l10n.ve.vat.book.export.wizard&id={self.id}'
                   f'&field=template_file&download=true&filename={filename}',
            'target': 'self',
        }

    def _generate_vat_book_xlsx(self, book_type):
        """Genera Excel para un tipo de libro (purchase o sale)."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'COMPRAS' if book_type == 'purchase' else 'VENTAS'

        # Estilos
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(
            start_color='2C3E50', end_color='2C3E50', fill_type='solid'
        )
        header_align = Alignment(horizontal='center', wrap_text=True)
        total_font = Font(bold=True)
        total_fill = PatternFill(
            start_color='E8E8E8', end_color='E8E8E8', fill_type='solid'
        )
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )
        number_format = '#,##0.00'

        self._write_sheet(ws, book_type, header_font, header_fill, header_align,
                         total_font, total_fill, thin_border, number_format, get_column_letter)

        # Guardar en bytes
        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def _generate_vat_book_both_xlsx(self):
        """Genera Excel con 2 hojas: COMPRAS y VENTAS."""
        import openpyxl

        wb = openpyxl.Workbook()

        # Estilos comunes
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(
            start_color='2C3E50', end_color='2C3E50', fill_type='solid'
        )
        header_align = Alignment(horizontal='center', wrap_text=True)
        total_font = Font(bold=True)
        total_fill = PatternFill(
            start_color='E8E8E8', end_color='E8E8E8', fill_type='solid'
        )
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )
        number_format = '#,##0.00'

        # Hoja COMPRAS
        ws_compras = wb.active
        ws_compras.title = 'COMPRAS'
        self._write_sheet(ws_compras, 'purchase', header_font, header_fill, header_align,
                         total_font, total_fill, thin_border, number_format, get_column_letter)

        # Hoja VENTAS
        ws_ventas = wb.create_sheet('VENTAS')
        self._write_sheet(ws_ventas, 'sale', header_font, header_fill, header_align,
                         total_font, total_fill, thin_border, number_format, get_column_letter)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def _write_sheet(self, ws, book_type, header_font, header_fill, header_align,
                     total_font, total_fill, thin_border, number_format, get_column_letter):
        """Escribe una hoja completa (headers + datos + total)."""
        headers = self._get_export_headers(book_type)

        # Headers
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col_idx)].width = 25

        # Datos
        domain = [
            ('period_month', '=', self.period_month),
            ('company_id', '=', self.company_id.id),
            ('book_type', '=', book_type),
        ]
        lines = self.env['l10n.ve.vat.book.line'].search(domain, order='partner_id, invoice_number, operation_code, retention_direction')

        row_idx = 2
        totals = {header: 0.0 for header in headers}

        for line in lines:
            for col_idx, header in enumerate(headers, 1):
                value = self._get_cell_value(line, header, book_type)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.number_format = number_format
                cell.border = thin_border
                if isinstance(value, (int, float)):
                    totals[header] += value
            row_idx += 1

        # Total general
        total_row = row_idx
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=total_row, column=col_idx)
            cell.border = thin_border
            cell.font = total_font
            cell.fill = total_fill
            if col_idx == 1:
                cell.value = 'TOTAL GENERAL'
                cell.alignment = Alignment(horizontal='center', wrap_text=True)
            elif header in totals:
                cell.value = totals[header]
                cell.number_format = number_format

    def _get_cell_value(self, line, header, book_type):
        """Obtiene el valor para una celda según el header y tipo de libro."""
        if book_type == 'purchase':
            return self._get_purchase_cell_value(line, header)
        else:
            return self._get_sale_cell_value(line, header)

    def _get_purchase_cell_value(self, line, header):
        """Valor para columnas de libro de compras."""
        if header == 'R.I.F.':
            return line.partner_id.vat or ''
        elif header == 'Nombre o Razon Social':
            return line.partner_id.name or ''
        elif header == 'Tipo Doc':
            return ''  # No existe en modelo
        elif header == 'Numero de Factura':
            return line.invoice_number or ''
        elif header == 'Numero de Control':
            return line.control_number or ''
        elif header == 'Fecha':
            return ''  # No existe campo fecha en modelo
        elif header == 'Base Alicuota General 16%':
            return line.base_general or 0.0
        elif header == 'I.V.A. Alicuota General 16%':
            return line.vat_general or 0.0
        elif header == 'Base Alicuota Reducida':
            return line.base_reduced or 0.0
        elif header == 'I.V.A. Alicuota Reducida':
            return line.vat_reduced or 0.0
        elif header == 'Compras NO SUJETAS':
            return line.base_not_subject or 0.0
        elif header == 'Compras sin Derecho a Credito (Nacional)':
            return line.base_no_credit or 0.0
        elif header == 'Base Importación 16%':
            return 0.0  # Se combina en base_general
        elif header == 'I.V.A. de importación 16%':
            return 0.0  # Se combina en vat_general
        elif header == 'Nº Comprob. Retención 75%':
            return line.retention_number or ''
        elif header == 'IVA Retenido (al Vendedor)':
            if line.retention_direction == 'to_vendor':
                return line.vat_retained or 0.0
            return 0.0
        elif header == 'IVA Retenido (a Terceros)':
            if line.retention_direction == 'to_third':
                return line.vat_retained or 0.0
            return 0.0
        elif header == 'Anticipo IVA (Importación)':
            return 0.0
        return 0.0

    def _get_sale_cell_value(self, line, header):
        """Valor para columnas de libro de ventas."""
        if header == 'R.I.F.':
            return line.partner_id.vat or ''
        elif header == 'Nombre o Razon Social':
            return line.partner_id.name or ''
        elif header == 'Tipo Doc':
            return ''  # No existe en modelo
        elif header == 'Numero de Factura':
            return line.invoice_number or ''
        elif header == 'Numero de Control':
            return line.control_number or ''
        elif header == 'Fecha':
            return ''  # No existe campo fecha en modelo
        elif header == 'Ventas internas No sujetas':
            return line.base_not_subject or 0.0
        elif header == 'Ventas internas no gravadas (No Contrib)':
            return line.base_not_taxed or 0.0
        elif header == 'Base Imponible (No Contrib)':
            # Para ventas no contrib, la base general corresponde a operation_code 40/41
            if line.operation_code in ('40', '41'):
                return line.base_general or 0.0
            return 0.0
        elif header == 'Impuesto IVA (No Contrib)':
            if line.operation_code in ('40', '41'):
                return line.vat_general or 0.0
            return 0.0
        elif header == 'Base Imponible (Contrib)':
            # Para ventas contrib (gravadas 16%), operation_code 42/442/443
            if line.operation_code in ('42', '442', '443'):
                return line.base_general or 0.0
            return 0.0
        elif header == 'Impuesto IVA (Contrib)':
            if line.operation_code in ('42', '442', '443'):
                return line.vat_general or 0.0
            return 0.0
        elif header == 'Nº Comprob. Retención 75% IVA':
            return line.retention_number or ''
        elif header == 'Iva Retenido (por comprador)':
            if line.retention_direction == 'by_buyer':
                return line.vat_retained or 0.0
            return 0.0
        return 0.0

    def _get_export_headers(self, book_type):
        """Retorna los headers en el orden exacto del template de importación."""
        if book_type == 'purchase':
            return [
                'R.I.F.',
                'Nombre o Razon Social',
                'Tipo Doc',
                'Numero de Factura',
                'Numero de Control',
                'Fecha',
                'Base Alicuota General 16%',
                'I.V.A. Alicuota General 16%',
                'Base Alicuota Reducida',
                'I.V.A. Alicuota Reducida',
                'Compras NO SUJETAS',
                'Compras sin Derecho a Credito (Nacional)',
                'Base Importación 16%',
                'I.V.A. de importación 16%',
                'Nº Comprob. Retención 75%',
                'IVA Retenido (al Vendedor)',
                'IVA Retenido (a Terceros)',
                'Anticipo IVA (Importación)',
            ]
        else:
            return [
                'R.I.F.',
                'Nombre o Razon Social',
                'Tipo Doc',
                'Numero de Factura',
                'Numero de Control',
                'Fecha',
                'Ventas internas No sujetas',
                'Ventas internas no gravadas (No Contrib)',
                'Base Imponible (No Contrib)',
                'Impuesto IVA (No Contrib)',
                'Base Imponible (Contrib)',
                'Impuesto IVA (Contrib)',
                'Nº Comprob. Retención 75% IVA',
                'Iva Retenido (por comprador)',
            ]
